from __future__ import annotations
from dataclasses import dataclass
from typing import List, Dict, Tuple, Iterable, Optional, Set
from datetime import date, datetime, timedelta
from src.db.sqlite import get_conn

@dataclass
class Slot:
    day_date: str   
    start_time: str 


def _to_minutes(d: date, time_str: str) -> int:
    """date + 'HH:MM' -> dakikaya çevir (epoch'a göre)."""
    hh, mm = [int(x) for x in time_str.split(":")]
    dt = datetime(d.year, d.month, d.day, hh, mm)
    return int(dt.timestamp() // 60)

def _duration_for(course_id: int, default_min: int, overrides: Optional[Dict[int, int]]) -> int:
    """
    Her ders için süreyi belirler:
      - overrides varsa ve ders burada listeliyse -> o süre
      - aksi halde varsayılan süre (default_min)
    """
    if overrides and course_id in overrides:
        return int(overrides[course_id])
    return int(default_min)

def build_slots(start_date: date, end_date: date,
                start_time: str, end_time: str,
                duration_min: int, gap_min: int,
                excluded_weekdays: Set[int]) -> List[Slot]:
    """
    start_time / end_time: 'HH:MM'
    duration_min + gap_min dikkate alınarak otomatik slot üretir.
    """
    slots: List[Slot] = []
    hh1, mm1 = map(int, start_time.split(":"))
    hh2, mm2 = map(int, end_time.split(":"))
    cur = start_date
    while cur <= end_date:
        if cur.weekday() not in (excluded_weekdays or set()):
            t = datetime(cur.year, cur.month, cur.day, hh1, mm1)
            t_end = datetime(cur.year, cur.month, cur.day, hh2, mm2)
            while (t + timedelta(minutes=duration_min)) <= t_end:
                slots.append(Slot(cur.isoformat(), t.strftime("%H:%M")))
                t += timedelta(minutes=duration_min + gap_min)
        cur += timedelta(days=1)
    return slots

def fetch_courses_with_counts(department_id: int | None = None,
                              include_ids: Optional[Iterable[int]] = None) -> List[dict]:
    con = get_conn(); cur = con.cursor()
    base = """
        SELECT c.id, c.code, c.name, COUNT(e.student_id) AS student_count
        FROM courses c
        LEFT JOIN enrollments e ON e.course_id = c.id
    """
    conds = []
    params: List[object] = []
    if department_id:
        conds.append("c.department_id = ?"); params.append(department_id)
    if include_ids:
        placeholders = ",".join("?" * len(list(include_ids)))
        conds.append(f"c.id IN ({placeholders})")
        params.extend(list(include_ids))
    where = (" WHERE " + " AND ".join(conds)) if conds else ""
    q = base + where + " GROUP BY c.id"
    cur.execute(q, params)
    rows = cur.fetchall(); con.close()
    return rows

def fetch_conflicts(course_ids: List[int]) -> Dict[int, set]:
    """Aynı öğrenciyi paylaşan dersleri bulur (çakışma grafı)."""
    if not course_ids:
        return {}
    con = get_conn(); cur = con.cursor()
    ph = ",".join("?" * len(course_ids))
    q = f"""
        SELECT e1.course_id AS c1, e2.course_id AS c2
        FROM enrollments e1
        JOIN enrollments e2 ON e1.student_id = e2.student_id
        WHERE e1.course_id != e2.course_id
          AND e1.course_id IN ({ph})
          AND e2.course_id IN ({ph})
    """
    cur.execute(q, course_ids + course_ids)
    adj = {cid: set() for cid in course_ids}
    for r in cur.fetchall():
        adj[r["c1"]].add(r["c2"]); adj[r["c2"]].add(r["c1"])
    con.close()
    return adj

def fetch_rooms(department_id: int | None = None, room_ids: Optional[Iterable[int]] = None) -> List[dict]:
    con = get_conn(); cur = con.cursor()
    base = "SELECT id, code, name, capacity FROM rooms"
    conds = []
    params: List[object] = []
    if department_id:
        conds.append("department_id = ?"); params.append(department_id)
    if room_ids:
        ids = list(room_ids)
        if ids:
            conds.append(f"id IN ({','.join('?'*len(ids))})")
            params.extend(ids)
    where = (" WHERE " + " AND ".join(conds)) if conds else ""
    q = base + where + " ORDER BY capacity DESC"
    cur.execute(q, params)
    rows = cur.fetchall(); con.close()
    return rows

def clear_existing_exams(exam_type: str, department_id: int | None) -> None:
    con = get_conn(); cur = con.cursor()
    if department_id:
        cur.execute("""
            DELETE FROM exam_rooms
             WHERE exam_id IN (
                SELECT ex.id FROM exams ex
                JOIN courses c ON c.id = ex.course_id
               WHERE ex.exam_type = ? AND c.department_id = ?
            )
        """, (exam_type, department_id))
        cur.execute("""
            DELETE FROM exams
             WHERE exam_type = ?
               AND course_id IN (SELECT id FROM courses WHERE department_id=?)
        """, (exam_type, department_id))
    else:
        cur.execute("DELETE FROM exam_rooms WHERE exam_id IN (SELECT id FROM exams WHERE exam_type=?)", (exam_type,))
        cur.execute("DELETE FROM exams WHERE exam_type=?", (exam_type,))
    con.commit(); con.close()

def schedule_exams(
    department_id: int | None,
    exam_type: str,
    start_date: date,
    end_date: date,
    *,
    start_time: str = "09:00",
    end_time: str = "17:00",
    default_duration_min: int = 75,
    excluded_weekdays: Optional[Set[int]] = None,
    min_gap_min: int = 15,
    single_at_a_time: bool = False,
    use_all_rooms: bool = True,
    room_ids: Optional[List[int]] = None,
    include_course_ids: Optional[List[int]] = None,
    duration_overrides: Optional[Dict[int, int]] = None,
) -> Tuple[int, List[str]]:
    """
    Otomatik sınav planlayıcı (greedy yaklaşım).
    - Günlük zaman aralığı (start_time–end_time) içinde slotları üretir.
    - Varsayılan sınav süresi default_duration_min, fakat
      duration_overrides sözlüğünde verilen dersler farklı sürelerle planlanır.
    - Çakışma, bekleme, global tek-sınav ve kapasite kısıtlarını dikkate alır.
    Döndürür: (yerleşen_sayısı, uyarılar_listesi)
    """
    warnings: List[str] = []
    excluded_weekdays = excluded_weekdays or set()


    courses = fetch_courses_with_counts(department_id, include_course_ids)
    courses.sort(key=lambda r: int(r["student_count"]), reverse=True)
    course_ids = [int(r["id"]) for r in courses]
    conflicts = fetch_conflicts(course_ids)
    rooms = fetch_rooms(department_id, room_ids if not use_all_rooms else None)
    if not rooms:
        return 0, ["Derslik bulunamadı."]

    slots = build_slots(
        start_date, end_date,
        start_time, end_time,
        default_duration_min, min_gap_min,
        excluded_weekdays
    )
    if not slots:
        return 0, ["Seçilen tarih aralığı/saatlere uygun slot yok."]

    assigned_at_slot: Dict[Tuple[str, str], List[int]] = {}
    room_busy: Dict[Tuple[str, str, int], bool] = {}
    student_last_end_min: Dict[int, int] = {}

    clear_existing_exams(exam_type, department_id)
    con = get_conn(); cur = con.cursor()
    placed = 0

    for c in courses:
        cid   = int(c["id"])
        need  = int(c["student_count"])
        ccode = c["code"]
        cname = c["name"]

        if need == 0:
            warnings.append(f"[{ccode}] için öğrenci yok.")
            continue

        cur.execute("SELECT student_id FROM enrollments WHERE course_id=?", (cid,))
        st_ids = [int(r["student_id"]) for r in cur.fetchall()]

        duration_min = _duration_for(cid, default_duration_min, duration_overrides)

        tried_any_slot = False
        ever_capacity_ok = False
        best_cap_slot = None
        best_cap_value = -1

        conflict_examples: List[str] = []
        gap_examples: List[str] = []
        global_block_examples: List[str] = []

        placed_this = False
        for sl in slots:
            tried_any_slot = True

            if single_at_a_time and assigned_at_slot.get((sl.day_date, sl.start_time)):
                global_block_examples.append(_fmt_slot(sl.day_date, sl.start_time))
                continue

            slot_courses = assigned_at_slot.get((sl.day_date, sl.start_time), [])
            slot_conflict_ids = [other for other in slot_courses if other in conflicts.get(cid, set())]
            if slot_conflict_ids:
                if len(conflict_examples) < 3:
                    id2code = {int(r["id"]): r["code"] for r in courses}
                    names = [id2code.get(int(x), str(x)) for x in slot_conflict_ids]
                    conflict_examples.append(f"{_fmt_slot(sl.day_date, sl.start_time)} -> {', '.join(names)}")
                continue

            d_obj = date.fromisoformat(sl.day_date)
            slot_start_min = _to_minutes(d_obj, sl.start_time)
            slot_end_min   = slot_start_min + duration_min
            if any((slot_start_min - student_last_end_min.get(sid, -10**12)) < min_gap_min for sid in st_ids):
                if len(gap_examples) < 3:
                    gap_examples.append(f"{_fmt_slot(sl.day_date, sl.start_time)} (min {min_gap_min} dk)")
                continue

            free_cap = _sum_capacity_free(rooms, room_busy, sl.day_date, sl.start_time)
            if free_cap > best_cap_value:
                best_cap_value = free_cap
                best_cap_slot = _fmt_slot(sl.day_date, sl.start_time)
            if free_cap < need:
                continue

            ever_capacity_ok = True

            remaining = need
            selected_rooms = []
            for r in rooms:
                if room_busy.get((sl.day_date, sl.start_time, int(r["id"])), False):
                    continue
                selected_rooms.append(r)
                remaining -= int(r["capacity"])
                if remaining <= 0:
                    break

            if remaining > 0:
                continue

            cur.execute("""
                INSERT INTO exams(course_id, exam_type, date, start_time, duration_min)
                VALUES(?,?,?,?,?)
            """, (cid, exam_type, sl.day_date, sl.start_time, duration_min))
            ex_id = cur.lastrowid

            for r in selected_rooms:
                cur.execute("INSERT INTO exam_rooms(exam_id, room_id) VALUES(?,?)", (ex_id, int(r["id"])) )
                room_busy[(sl.day_date, sl.start_time, int(r["id"]))] = True

            con.commit()

            assigned_at_slot.setdefault((sl.day_date, sl.start_time), []).append(cid)
            for sid in st_ids:
                student_last_end_min[sid] = slot_end_min

            placed += 1
            placed_this = True
            break

        if not placed_this:
            if not tried_any_slot:
                warnings.append(f"[{ccode}] ({cname}) için slot yok (tarih aralığı/hariç günler tümünü kesti).")
            else:
                if not ever_capacity_ok:
                    cap_info = f"ihtiyaç {need}, en iyi slot kapasite {max(0, best_cap_value)}"
                    if best_cap_slot:
                        warnings.append(f"[{ccode}] ({cname}) kapasite yetersiz: {cap_info} @ {best_cap_slot}.")
                    else:
                        warnings.append(f"[{ccode}] ({cname}) kapasite yetersiz: {cap_info}.")
                elif conflict_examples:
                    warnings.append(f"[{ccode}] ({cname}) çakışma: " + "; ".join(conflict_examples) + ".")
                elif gap_examples:
                    warnings.append(f"[{ccode}] ({cname}) bekleme kısıtı: " + "; ".join(gap_examples) + ".")
                elif global_block_examples:
                    warnings.append(f"[{ccode}] ({cname}) global 'aynı anda tek sınav' nedeniyle yer bulamadı; "
                                    f"örnek: {', '.join(global_block_examples[:3])}.")
                else:
                    warnings.append(f"[{ccode}] ({cname}) yerleştirilemedi (kısıtlar nedeniyle uygun slot/oda yok).")

    con.close()
    return placed, warnings


def _fmt_slot(day_date: str, start_time: str) -> str:
    return f"{day_date} {start_time}"

def _sum_capacity_free(rooms: List[dict], room_busy: dict, day_date: str, start_time: str) -> int:
    total = 0
    for r in rooms:
        if not room_busy.get((day_date, start_time, int(r["id"])), False):
            total += int(r["capacity"])
    return total


def list_scheduled(exam_type: str, department_id: int | None = None) -> List[dict]:
    con = get_conn(); cur = con.cursor()
    if department_id:
        cur.execute("""
            SELECT ex.id, c.code, c.name, ex.date, ex.start_time, ex.duration_min, r.code AS room_code
              FROM exams ex
              JOIN courses c ON c.id = ex.course_id
              JOIN exam_rooms er ON er.exam_id = ex.id
              JOIN rooms r ON r.id = er.room_id
             WHERE ex.exam_type = ? AND c.department_id = ?
          ORDER BY ex.date, ex.start_time, c.code
        """, (exam_type, department_id))
    else:
        cur.execute("""
            SELECT ex.id, c.code, c.name, ex.date, ex.start_time, ex.duration_min, r.code AS room_code
              FROM exams ex
              JOIN courses c ON c.id = ex.course_id
              JOIN exam_rooms er ON er.exam_id = ex.id
              JOIN rooms r ON r.id = er.room_id
             WHERE ex.exam_type = ?
          ORDER BY ex.date, ex.start_time, c.code
        """, (exam_type,))
    rows = cur.fetchall(); con.close()
    return rows

def export_schedule(exam_type: str, department_id: int | None, path: str) -> str:
    """
    Verilen exam_type (+ opsiyonel department_id) için planlanan sınavları
    path'e yazar. .xlsx ise openpyxl ile Excel, aksi halde CSV üretir.
    Dönen değer: kaydedilen dosya yolu.
    """
    rows = list_scheduled(exam_type, department_id)
    headers = ["Ders Kodu", "Ders Adı", "Tarih", "Saat", "Süre (dk)", "Derslik"]

    if path.lower().endswith(".xlsx"):
        try:
            from openpyxl import Workbook
        except ImportError as e:
            raise RuntimeError("Excel çıktısı için 'openpyxl' gerekiyor. "
                               "İsterseniz CSV olarak kaydedin veya 'pip install openpyxl' kurun.") from e

        wb = Workbook()
        ws = wb.active
        ws.title = "Sınav Programı"
        ws.append(headers)
        for r in rows:
            ws.append([r["code"], r["name"], r["date"], r["start_time"], int(r["duration_min"]), r["room_code"]])

        for col in ws.columns:
            max_len = 0
            col_letter = col[0].column_letter
            for cell in col:
                val = "" if cell.value is None else str(cell.value)
                if len(val) > max_len:
                    max_len = len(val)
            ws.column_dimensions[col_letter].width = max(12, min(40, max_len + 2))

        wb.save(path)
        return path

    import csv
    with open(path, "w", newline="", encoding="utf-8") as f:
        w = csv.writer(f)
        w.writerow(headers)
        for r in rows:
            w.writerow([r["code"], r["name"], r["date"], r["start_time"], int(r["duration_min"]), r["room_code"]])
    return path
